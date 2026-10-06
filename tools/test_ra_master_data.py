"""Data rules for the Master Rate Analysis block (#raMasterData) and the Rev07 workbook.

    python3 -m unittest tools/test_ra_master_data.py
"""
import datetime as dt
import json
import re
import sys
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ra_master_config as C  # noqa: E402
from ra_master_data import dash_lines, seed_history  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
HTML = (ROOT / "zameen-developments" / "index.html").read_text(encoding="utf-8")
BLOCK = json.loads(re.search(r'id="raMasterData">(.*?)</script>', HTML, re.S).group(1))
WB = ROOT / "docs" / "RA_Master_Pakistan_FINAL_Rev07_2026-09-27.xlsx"
DATED = re.compile(r"\d{2}-[A-Z][a-z]{2}-\d{4}")


class Block(unittest.TestCase):
    def test_every_line_sourced(self):
        for r in BLOCK["rates"]:
            ok = r["src"].startswith(("ASSUMPTION", "ASSUMPTION —")) or DATED.search(r["src"])
            self.assertTrue(ok, r["code"])
            if r["rate"] and not r["src"].startswith("ASSUMPTION"):
                self.assertTrue(r["date"], r["code"])
            if not r["date"]:
                self.assertTrue(r["src"].startswith("ASSUMPTION"), r["code"])

    def test_remark_date_is_effective_date(self):
        """CLAUDE.md form `<source>, DD-Mon-YYYY — <details>`: one date before the dash, equal to the line's
        effective date (a price list is dated by its effective date, not the day it was read)."""
        for r in BLOCK["rates"]:
            if r["src"].startswith("ASSUMPTION") or not r["date"]:
                continue
            self.assertIn(" — ", r["src"], r["code"])
            head = DATED.findall(r["src"].split(" — ")[0])
            self.assertEqual(len(head), 1, f"{r['code']}: {r['src'][:90]}")
            self.assertEqual(dt.datetime.strptime(head[0], "%d-%b-%Y").date().isoformat(), r["date"], r["code"])

    def test_workbook_price_list_dates(self):
        w05 = openpyxl.load_workbook(WB, data_only=True)["05 MATERIAL RATE LIBRARY"]
        for r in range(5, w05.max_row + 1):
            if str(w05.cell(r, 11).value or "") in ("SRC-25", "SRC-26"):  # Popular Pipes conduit / Fast Cables
                self.assertLessEqual(str(w05.cell(r, 9).value), "2026-09-14", w05.cell(r, 1).value)

    def test_rows_point_at_lines(self):
        codes = {r["code"] for r in BLOCK["rates"]} | set(dash_lines(HTML))
        for it in BLOCK["items"]:
            for k in "MLP":
                for row in it[k]:
                    self.assertIn(row["ref"], codes, it["id"])

    def test_items_carry_categories_and_no_duplicate(self):
        ids = [it["id"] for it in BLOCK["items"]]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertGreaterEqual(len(ids), 160)
        for it in BLOCK["items"]:
            self.assertRegex(it["qc"], r"^C\d\d$")
            self.assertEqual((it["wast"], it["oh"], it["prof"]), (0, 8, 10))

    def test_shared_lines_not_redefined(self):
        mine = {r["code"] for r in BLOCK["rates"]}
        for x, d in C.MAP.items():
            self.assertNotIn(x, mine)
            self.assertTrue(d not in mine or d in C.DASH_RATE_UPDATES, d)


@unittest.skipUnless(WB.exists(), "Rev07 workbook not in docs/")
class Workbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.wv = openpyxl.load_workbook(WB, data_only=True)

    def test_shared_rates_equal(self):
        cur = dash_lines(HTML)
        cur.update({k: dict(cur[k], **{a: b for a, b in v.items() if a != "prev"})
                    for k, v in C.DASH_RATE_UPDATES.items()})
        rates = {}
        for s, col in (("05 MATERIAL RATE LIBRARY", 5), ("06 LABOUR RATE LIBRARY", 4), ("07 EQUIPMENT RATE LIBRARY", 4)):
            w = self.wv[s]
            for r in range(5, w.max_row + 1):
                rates[w.cell(r, 1).value] = w.cell(r, col).value
        hist = seed_history(HTML)   # a line updated after the workbook (newer GRN) matches through its history
        for x, d in C.MAP.items():
            ok = {float(cur[d]["rate"])} | hist.get(d, set())
            self.assertTrue(any(abs(float(rates[x]) - v) < 1e-6 for v in ok), f"{x} {rates[x]} vs {d} {sorted(ok)}")

    def test_reconciliation_sheet_all_ok(self):
        w = self.wv["18 DASHBOARD RECONCILIATION"]
        res = [w.cell(r, 14).value for r in range(5, w.max_row + 1) if w.cell(r, 1).value and
               str(w.cell(r, 1).value).count("-") >= 1 and w.cell(r, 14).value in ("OK",) or
               str(w.cell(r, 14).value or "").startswith("CHECK")]
        self.assertTrue(res and all(v == "OK" for v in res))

    def test_items_match_block(self):
        w = self.wv["03 MASTER RATE LIBRARY"]
        codes = {w.cell(r, 1).value for r in range(5, w.max_row + 1) if w.cell(r, 1).value}
        self.assertEqual(codes, {it["id"] for it in BLOCK["items"]})

    def test_no_formula_errors(self):
        for ws in self.wv:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str):
                        self.assertNotRegex(c.value, r"^#(REF|VALUE|N/A|DIV/0|NAME)", f"{ws.title}!{c.coordinate}")


if __name__ == "__main__":
    unittest.main()
